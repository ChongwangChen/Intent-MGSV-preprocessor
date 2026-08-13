from __future__ import annotations

import argparse
import json
import os
import re
import socket
import subprocess
import sys
import threading
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable, TextIO, TypeVar
from urllib.parse import urlparse

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DOUK_ROOT = PROJECT_ROOT / "DouK-Source"
DEFAULT_METADATA = DEFAULT_DOUK_ROOT / "Volume" / "Data" / "Download.xlsx"
DEFAULT_OUTPUT = PROJECT_ROOT / "outputs" / "server" / "douyin_batch_download"
WORK_ID_PATTERN = re.compile(r"/(?:video|note)/(\d+)")
T = TypeVar("T")


@dataclass(frozen=True)
class WorkLink:
    work_id: str
    url: str


def extract_work_id(url: str) -> str | None:
    match = WORK_ID_PATTERN.search(url)
    return match.group(1) if match else None


def read_links(package_dir: Path) -> tuple[list[WorkLink], int, int]:
    merged = package_dir / "douyin_links_merged.txt"
    sources = [merged] if merged.is_file() else sorted(package_dir.glob("dydownload_*.txt"))
    if not sources:
        raise FileNotFoundError(
            f"No douyin_links_merged.txt or dydownload_*.txt found in {package_dir}"
        )

    seen_ids: set[str] = set()
    links: list[WorkLink] = []
    malformed = 0
    duplicate_ids = 0
    for source in sources:
        for raw in source.read_text(encoding="utf-8-sig").splitlines():
            url = raw.strip()
            if not url:
                continue
            work_id = extract_work_id(url)
            if not work_id:
                malformed += 1
                continue
            if work_id in seen_ids:
                duplicate_ids += 1
                continue
            seen_ids.add(work_id)
            links.append(WorkLink(work_id=work_id, url=url))
    return links, malformed, duplicate_ids


def read_original_batch_files(package_dir: Path) -> list[Path]:
    batches = sorted(package_dir.glob("dydownload_*.txt"))
    if not batches:
        raise FileNotFoundError(f"No dydownload_*.txt found in {package_dir}")
    return [path for path in batches if path.stat().st_size > 0]


def _normalize_header(value: object) -> str:
    return str(value or "").strip().lower().replace("_", "")


def _normalize_work_id(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value or "").strip()


def read_downloaded_ids(metadata_path: Path) -> set[str]:
    if not metadata_path.is_file():
        return set()

    last_error: Exception | None = None
    for attempt in range(3):
        try:
            workbook = load_workbook(metadata_path, read_only=True, data_only=True)
            try:
                sheet = workbook.active
                rows = sheet.iter_rows(values_only=True)
                headers = next(rows, ())
                normalized = [_normalize_header(value) for value in headers]
                id_index = None
                for candidate in ("作品ID", "作品id", "id"):
                    normalized_candidate = _normalize_header(candidate)
                    if normalized_candidate in normalized:
                        id_index = normalized.index(normalized_candidate)
                        break
                if id_index is None:
                    raise ValueError(f"No work ID column found in {metadata_path}")
                result = {
                    _normalize_work_id(row[id_index])
                    for row in rows
                    if id_index < len(row) and row[id_index] not in (None, "")
                }
                return result
            finally:
                workbook.close()
        except (OSError, ValueError) as exc:
            last_error = exc
            time.sleep(attempt + 1)
    raise RuntimeError(f"Unable to read DouK metadata: {last_error}")


def chunked(items: list[T], size: int) -> Iterable[list[T]]:
    for start in range(0, len(items), size):
        yield items[start : start + size]


def write_batch(path: Path, links: list[WorkLink]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(item.url for item in links) + "\n", encoding="utf-8")


def configured_proxy(settings_path: Path) -> str:
    if not settings_path.is_file():
        return ""
    try:
        data = json.loads(settings_path.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError):
        return ""
    return str(data.get("proxy", "") or "").strip()


def check_proxy_listener(proxy: str) -> tuple[bool, str]:
    if not proxy:
        return True, "DouK proxy is disabled"
    parsed = urlparse(proxy)
    if not parsed.hostname or parsed.port is None:
        return False, f"Invalid proxy setting: {proxy}"
    try:
        with socket.create_connection((parsed.hostname, parsed.port), timeout=3):
            pass
    except OSError as exc:
        return False, f"Proxy listener unavailable at {parsed.hostname}:{parsed.port}: {exc}"
    return True, f"Proxy listener reachable: {proxy}"


def _tee_output(stream: TextIO, log_handle: TextIO) -> None:
    for line in iter(stream.readline, ""):
        print(line, end="", flush=True)
        log_handle.write(line)
        log_handle.flush()


def run_douk_batch(
    *,
    python: str,
    douk_root: Path,
    batch_path: Path,
    log_path: Path,
    timeout_seconds: int,
) -> tuple[int, str]:
    return run_douk_batch_group(
        python=python,
        douk_root=douk_root,
        batch_paths=[batch_path],
        log_path=log_path,
        timeout_seconds=timeout_seconds,
    )


def build_douk_commands(batch_paths: list[Path]) -> str:
    commands = ["3"]
    for batch_path in batch_paths:
        commands.extend(("2", "2", str(batch_path.resolve())))
    commands.append("q")
    return "\n".join(commands) + "\n"


def run_douk_batch_group(
    *,
    python: str,
    douk_root: Path,
    batch_paths: list[Path],
    log_path: Path,
    timeout_seconds: int,
) -> tuple[int, str]:
    if not batch_paths:
        raise ValueError("batch_paths must not be empty")
    commands = build_douk_commands(batch_paths)
    env = os.environ.copy()
    env["PYTHONUNBUFFERED"] = "1"
    log_path.parent.mkdir(parents=True, exist_ok=True)

    with log_path.open("w", encoding="utf-8", errors="replace") as log_handle:
        process = subprocess.Popen(
            [python, "-u", str(douk_root / "main.py")],
            cwd=douk_root,
            env=env,
            stdin=subprocess.PIPE,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            encoding="utf-8",
            errors="replace",
            bufsize=1,
        )
        assert process.stdin is not None
        assert process.stdout is not None
        process.stdin.write(commands)
        process.stdin.close()
        reader = threading.Thread(
            target=_tee_output,
            args=(process.stdout, log_handle),
            daemon=True,
        )
        reader.start()
        try:
            return_code = process.wait(timeout=timeout_seconds)
            reason = "completed"
        except subprocess.TimeoutExpired:
            process.terminate()
            try:
                process.wait(timeout=20)
            except subprocess.TimeoutExpired:
                process.kill()
                process.wait()
            return_code = 124
            reason = f"timed out after {timeout_seconds}s"
        reader.join(timeout=10)
    return return_code, reason


def summarize_log(log_path: Path) -> dict[str, int]:
    text = log_path.read_text(encoding="utf-8", errors="replace")
    lines = text.splitlines()
    return {
        "video_download_success": sum(
            ("【视频】" in line) and ("文件下载成功" in line) for line in lines
        ),
        "gallery_download_success": sum(
            ("【图集】" in line) and ("文件下载成功" in line) for line in lines
        ),
        "music_download_success": sum(
            (("【Music】" in line) or ("【音乐】" in line))
            and ("文件下载成功" in line)
            for line in lines
        ),
        "existing_skipped": sum(
            ("存在下载记录" in line) or ("文件已存在" in line) for line in lines
        ),
        "download_interrupted": sum("下载中断" in line for line in lines),
        "detail_fetch_failed": sum("获取作品数据失败" in line for line in lines),
        "http_403": sum("403 Forbidden" in line for line in lines),
        "tls_eof": sum(
            ("UNEXPECTED_EOF_WHILE_READING" in line)
            or ("unexpected eof while reading" in line.lower())
            for line in lines
        ),
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Filter links already present in DouK Download.xlsx, then drive "
            "DouK main.py through all remaining batches without menu input."
        )
    )
    parser.add_argument("package_dir", help="Directory containing merged/batch link TXT files")
    parser.add_argument("--douk-root", default=str(DEFAULT_DOUK_ROOT))
    parser.add_argument("--metadata", default=str(DEFAULT_METADATA))
    parser.add_argument("--output-dir", default=str(DEFAULT_OUTPUT))
    parser.add_argument("--batch-size", type=int, default=20)
    parser.add_argument(
        "--batches-per-process",
        type=int,
        default=4,
        help=(
            "Feed this many small batches to one DouK process. This avoids "
            "repeating DouK startup and parameter refresh for every batch."
        ),
    )
    parser.add_argument("--max-rounds", type=int, default=2)
    parser.add_argument("--pause-seconds", type=int, default=30)
    parser.add_argument("--batch-timeout-seconds", type=int, default=3600)
    parser.add_argument(
        "--abort-after-network-errors",
        type=int,
        default=10,
        help=(
            "Stop before later batches when one batch completes no new work IDs "
            "and reaches this many TLS/detail/403 errors; 0 disables the fuse."
        ),
    )
    parser.add_argument("--python", default=sys.executable)
    parser.add_argument(
        "--run-original-batches",
        action="store_true",
        help=(
            "Run every dydownload_*.txt exactly once, including works already "
            "recorded by DouK, so missing music/media can be supplemented."
        ),
    )
    parser.add_argument("--dry-run", action="store_true")
    return parser


def main() -> None:
    args = build_parser().parse_args()
    if (
        args.batch_size <= 0
        or args.batches_per_process <= 0
        or args.max_rounds <= 0
    ):
        raise SystemExit(
            "--batch-size, --batches-per-process, and --max-rounds must be positive"
        )

    package_dir = Path(args.package_dir).expanduser().resolve()
    douk_root = Path(args.douk_root).expanduser().resolve()
    metadata = Path(args.metadata).expanduser().resolve()
    settings = douk_root / "Volume" / "settings.json"
    session_dir = (
        Path(args.output_dir).expanduser().resolve()
        / datetime.now().strftime("%Y%m%d_%H%M%S")
    )
    session_dir.mkdir(parents=True, exist_ok=True)

    links, malformed, duplicate_ids = read_links(package_dir)
    original_batches = read_original_batch_files(package_dir)
    downloaded_before = read_downloaded_ids(metadata)
    pending = [item for item in links if item.work_id not in downloaded_before]
    proxy = configured_proxy(settings)
    proxy_ok, proxy_message = check_proxy_listener(proxy)

    initial = {
        "package_dir": str(package_dir),
        "source_unique_links": len(links),
        "malformed_links_skipped": malformed,
        "duplicate_ids_removed": duplicate_ids,
        "already_downloaded_removed": len(links) - len(pending),
        "pending": len(pending),
        "metadata": str(metadata),
        "proxy": proxy,
        "proxy_check": proxy_message,
        "session_dir": str(session_dir),
        "execution_mode": (
            "original_batches" if args.run_original_batches else "pending_work_ids"
        ),
        "batches_per_process": args.batches_per_process,
        "original_batch_files": [path.name for path in original_batches],
    }
    print(json.dumps(initial, ensure_ascii=False, indent=2))
    (session_dir / "initial.json").write_text(
        json.dumps(initial, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    if proxy and not proxy_ok:
        raise SystemExit(proxy_message)

    if args.dry_run:
        if pending:
            write_batch(session_dir / "pending_dry_run.txt", pending)
        return


    if args.run_original_batches:
        batch_runs: list[dict[str, object]] = []
        process_groups = list(
            chunked(original_batches, args.batches_per_process)
        )
        for group_number, batch_paths in enumerate(process_groups, start=1):
            log_path = (
                session_dir
                / "original_batches"
                / f"douk_process_{group_number:02d}.log"
            )
            submitted = sum(
                sum(
                    bool(line.strip())
                    for line in batch_path.read_text(
                        encoding="utf-8-sig"
                    ).splitlines()
                )
                for batch_path in batch_paths
            )
            print(
                f"\n[{group_number}/{len(process_groups)}] Running one DouK "
                f"process for {len(batch_paths)} original batches / "
                f"{submitted} links"
            )
            return_code, reason = run_douk_batch_group(
                python=args.python,
                douk_root=douk_root,
                batch_paths=batch_paths,
                log_path=log_path,
                timeout_seconds=args.batch_timeout_seconds,
            )
            result = {
                "process_group": group_number,
                "source_batches": [path.name for path in batch_paths],
                "submitted": submitted,
                "return_code": return_code,
                "reason": reason,
                "log_path": str(log_path),
                **summarize_log(log_path),
            }
            batch_runs.append(result)
            print(json.dumps(result, ensure_ascii=False))
            if args.pause_seconds > 0 and group_number < len(process_groups):
                time.sleep(args.pause_seconds)

        downloaded_final = read_downloaded_ids(metadata)
        totals = {
            key: sum(int(run[key]) for run in batch_runs)
            for key in (
                "video_download_success",
                "gallery_download_success",
                "music_download_success",
                "existing_skipped",
                "download_interrupted",
                "detail_fetch_failed",
                "http_403",
                "tls_eof",
            )
        }
        report = {
            **initial,
            "original_batches_executed": len(batch_runs),
            "new_work_ids": len(downloaded_final - downloaded_before),
            "log_totals": totals,
            "batch_runs": batch_runs,
        }
        report_path = session_dir / "report.json"
        report_path.write_text(
            json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        print("\nFinal report:")
        print(json.dumps(report, ensure_ascii=False, indent=2))
        print(f"Report: {report_path}")
        return

    if not pending:
        return

    batch_runs: list[dict[str, object]] = []
    network_aborted = False
    for round_number in range(1, args.max_rounds + 1):
        downloaded = read_downloaded_ids(metadata)
        pending = [item for item in links if item.work_id not in downloaded]
        if not pending:
            break
        print(f"\nRound {round_number}: {len(pending)} pending links")
        round_dir = session_dir / f"round_{round_number:02d}"
        batches = list(chunked(pending, args.batch_size))
        batch_specs: list[tuple[int, list[WorkLink], Path]] = []
        for batch_number, batch in enumerate(batches, start=1):
            batch_path = round_dir / f"dydownload_{batch_number:02d}.txt"
            write_batch(batch_path, batch)
            batch_specs.append((batch_number, batch, batch_path))

        process_groups = list(
            chunked(batch_specs, args.batches_per_process)
        )
        for group_number, group in enumerate(process_groups, start=1):
            first_batch = group[0][0]
            last_batch = group[-1][0]
            group_links = [item for _, batch, _ in group for item in batch]
            batch_paths = [batch_path for _, _, batch_path in group]
            log_path = round_dir / f"douk_process_{group_number:02d}.log"
            print(
                f"\n[{group_number}/{len(process_groups)}] Running one DouK "
                f"process for batches {first_batch}-{last_batch} "
                f"({len(group_links)} links)"
            )
            return_code, reason = run_douk_batch_group(
                python=args.python,
                douk_root=douk_root,
                batch_paths=batch_paths,
                log_path=log_path,
                timeout_seconds=args.batch_timeout_seconds,
            )
            downloaded_after = read_downloaded_ids(metadata)
            completed = sum(
                item.work_id in downloaded_after for item in group_links
            )
            log_summary = summarize_log(log_path)
            result = {
                "round": round_number,
                "process_group": group_number,
                "batch_start": first_batch,
                "batch_end": last_batch,
                "batch_paths": [str(path) for path in batch_paths],
                "submitted": len(group_links),
                "completed_after_group": completed,
                "return_code": return_code,
                "reason": reason,
                "log_path": str(log_path),
                **log_summary,
            }
            batch_runs.append(result)
            print(json.dumps(result, ensure_ascii=False))
            network_errors = (
                log_summary["detail_fetch_failed"]
                + log_summary["http_403"]
                + log_summary["tls_eof"]
            )
            if (
                args.abort_after_network_errors > 0
                and completed == 0
                and network_errors >= args.abort_after_network_errors
            ):
                network_aborted = True
                print(
                    "Network fuse opened: no work ID completed and "
                    f"network_errors={network_errors}. Stop and rebuild the proxy "
                    "before resuming unresolved_links.txt."
                )
                break
            if args.pause_seconds > 0 and group_number < len(process_groups):
                time.sleep(args.pause_seconds)

        if network_aborted:
            break
        if round_number < args.max_rounds and args.pause_seconds > 0:
            print(f"Waiting {args.pause_seconds}s before retrying unresolved links...")
            time.sleep(args.pause_seconds)

    downloaded_final = read_downloaded_ids(metadata)
    unresolved = [item for item in links if item.work_id not in downloaded_final]
    if unresolved:
        write_batch(session_dir / "unresolved_links.txt", unresolved)
    report = {
        **initial,
        "downloaded_during_run": len(downloaded_final - downloaded_before),
        "remaining_unresolved": len(unresolved),
        "network_aborted": network_aborted,
        "batch_runs": batch_runs,
        "unresolved_file": (
            str(session_dir / "unresolved_links.txt") if unresolved else ""
        ),
    }
    report_path = session_dir / "report.json"
    report_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print("\nFinal report:")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    print(f"Report: {report_path}")


if __name__ == "__main__":
    main()
